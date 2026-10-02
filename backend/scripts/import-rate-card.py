#!/usr/bin/env python3
"""
Target Logistics - Rate Card Importer CLI
Imports a customer/carrier zone pricing matrix Excel sheet (.xlsx) into a production JSON Rate Card.

Usage:
    python backend/scripts/import-rate-card.py --file "path/to/matrix.xlsx" --id "CLIENT_CODE" --name "Client Display Name" [--carrier DGR] [--currency KWD]

Example:
    python backend/scripts/import-rate-card.py --file "docs/Zone Pricing/5535 - Amani.xlsx" --id "5535_AMANI" --name "5535 - Amani"
"""

import argparse
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Run: pip install openpyxl")
    sys.exit(1)

def parse_args():
    parser = argparse.ArgumentParser(description="Import Excel rate card to Target Logistics JSON format.")
    parser.add_argument("--file", required=True, help="Path to Excel (.xlsx) file")
    parser.add_argument("--id", required=True, help="Unique identifier, e.g. 6000_KHALID")
    parser.add_argument("--name", required=True, help="Display name, e.g. '6000 - Khalid'")
    parser.add_argument("--carrier", default="DGR", help="Carrier code (default: DGR)")
    parser.add_argument("--currency", default="KWD", help="Currency code (default: KWD)")
    parser.add_argument("--mode", default="SELLING_PRICE", choices=["SELLING_PRICE", "BASE_COST"], help="Pricing mode")
    parser.add_argument("--sheet", default=None, help="Sheet name (default: first sheet)")
    return parser.parse_args()

def main():
    args = parse_args()
    file_path = os.path.abspath(args.file)
    if not os.path.exists(file_path):
        print(f"Error: File not found at {file_path}")
        sys.exit(1)

    print(f"Loading workbook: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active

    # Detect header columns
    headers = [str(ws.cell(1, c).value or '').strip() for c in range(1, ws.max_column + 1)]
    print(f"Detected columns: {headers}")

    weight_col = 1
    zone_cols = {}
    for idx, h in enumerate(headers):
        if idx == 0:
            continue
        cleaned = h.replace(' ', '').replace('.', '').upper()
        if 'ZONE' in cleaned:
            zone_num = cleaned.replace('ZONE', '')
            zone_cols[zone_num] = idx + 1

    print(f"Detected zones: {list(zone_cols.keys())}")

    brackets = []
    for r in range(2, ws.max_row + 1):
        raw_w = ws.cell(r, weight_col).value
        if raw_w is None:
            continue
        try:
            w = float(raw_w)
        except ValueError:
            continue

        rates = {}
        for z_num, col_idx in zone_cols.items():
            val = ws.cell(r, col_idx).value
            if val is not None:
                rates[str(z_num)] = round(float(val), 4)

        brackets.append({
            "weight": round(w, 2),
            "rates": rates
        })

    if not brackets:
        print("Error: No rate rows found!")
        sys.exit(1)

    brackets.sort(key=lambda x: x["weight"])
    max_weight = brackets[-1]["weight"]
    step = round(brackets[1]["weight"] - brackets[0]["weight"], 2) if len(brackets) > 1 else 0.5
    print(f"Imported {len(brackets)} brackets (from {brackets[0]['weight']} to {max_weight} kg, step: {step} kg).")

    # Compute marginal excess rate from last 2 brackets
    over30KgPerKgRate = {}
    if len(brackets) >= 2:
        last = brackets[-1]["rates"]
        second_last = brackets[-2]["rates"]
        w_diff = brackets[-1]["weight"] - brackets[-2]["weight"]
        for z_num in zone_cols.keys():
            if str(z_num) in last and str(z_num) in second_last and w_diff > 0:
                diff = last[str(z_num)] - second_last[str(z_num)]
                # Convert step diff to per-1.0kg rate
                per_kg = round(diff / w_diff, 4)
                over30KgPerKgRate[str(z_num)] = per_kg

    card_id = args.id.strip().upper()
    rate_card_data = {
        "id": card_id,
        "name": args.name.strip(),
        "carrierCode": args.carrier.strip().upper(),
        "currency": args.currency.strip().upper(),
        "pricingMode": args.mode,
        "maxBracketWeight": max_weight,
        "weightStep": step,
        "brackets": brackets,
        "over30KgPerKgRate": over30KgPerKgRate
    }

    # Output directory
    script_dir = os.path.dirname(os.path.abspath(__file__))
    target_dir = os.path.abspath(os.path.join(script_dir, "../src/constants/rateCards"))
    os.makedirs(target_dir, exist_ok=True)

    filename = f"{args.carrier.lower()}_{card_id.lower()}.json"
    target_path = os.path.join(target_dir, filename)

    with open(target_path, "w", encoding="utf-8") as f:
        json.dump(rate_card_data, f, indent=2, ensure_ascii=False)

    print(f"\nSuccessfully created Rate Card:")
    print(f"  ID: {card_id}")
    print(f"  Name: {args.name}")
    print(f"  Path: {target_path}")
    print(f"  Available in system automatically!")

if __name__ == "__main__":
    main()
